"""Genera la plantilla del Excel fuente (`plantilla_clarios.xlsx`) que descarga el front.

Las hojas de ejemplo empiezan por "_" para que el pipeline las ignore si se sube la plantilla tal cual.
Las reglas que se explican en la hoja Instrucciones son las de pipeline/esquema.py.

Uso: python -m pipeline.plantilla [ruta_salida]
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .config import COLUMNS_ENERGIA, COLUMNS_GAS, ROOT

SALIDA = ROOT / "entrada" / "plantilla_clarios.xlsx"
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

AZUL = PatternFill("solid", fgColor="1F3A5F")
GRIS = PatternFill("solid", fgColor="E8EDF3")
BLANCO_NEGRITA = Font(bold=True, color="FFFFFF")
NOTA = Font(italic=True, color="6B7280", size=9)
NEGRITA = Font(bold=True)

# Encabezado -> (nota de la fila 1, obligatoria, descripción para la hoja Instrucciones)
DESCRIPCION = {
    "Año": ("", True, "Año del mes, número (2026)."),
    "Mes": ("Ene..Dic", True, "Mes abreviado: Ene, Feb, ... Dic (también se acepta Jan..Dec)."),
    "Dia": ("N/A", False, "No se usa (datos mensuales). Puede ir N/A."),
    "Turno": ("N/A", False, "No se usa (datos mensuales). Puede ir N/A."),
    "Proceso": ("Último = total", True, "Nombre del proceso. La última fila de cada mes es el total y se llama "
                                        "como la planta."),
    "Indice": ("Suma 1 por mes", True, "Participación del proceso en el consumo de la planta (0 a 1). "
                                       "En la fila total va 1. También se acepta el encabezado 'Participacion'."),
    "KWH": ("SCADA", True, "Consumo de toda la planta en el mes (se repite en cada fila del mes)."),
    "KWH por proceso": ("KWH × Indice", True, "Consumo asignado al proceso. En la fila total es igual a KWH."),
    "(m3)": ("SCADA", True, "Consumo de gas de la planta en el mes."),
    "Produccion (Equivalente)": ("ECU", True, "Producción en unidades equivalentes. Vacío o 0 si no aplica."),
    "KWH/Equ (Reales)": ("", False, "Intensidad del Excel; el pipeline la recalcula."),
    "m3 /Equ (Reales)": ("", False, "Intensidad del Excel; el pipeline la recalcula."),
    "KWH/Equ (Presupuesto)": ("", False, "Intensidad presupuestada (meta). Vacío si no hay."),
    "m3 /Equ (presupuesto)": ("", False, "Intensidad presupuestada (meta). Vacío si no hay."),
    "Costo por KWH": ("COP", False, "Tarifa del mes. Vacío = la planta queda sin pronóstico de costo."),
    "Costo m3": ("COP", False, "Tarifa del mes. Vacío = la planta queda sin pronóstico de costo."),
    "Costo  total KWH": ("", False, "Costo del Excel; el pipeline lo recalcula."),
    "Total Costo m3": ("", False, "Costo del Excel; el pipeline lo recalcula."),
}

ENCABEZADOS = {
    "energia": [k for k in COLUMNS_ENERGIA if k != "Participacion"],
    "gas": list(COLUMNS_GAS),
}

REGLAS = [
    "Una hoja por planta y servicio. Nombre: '<Planta>' o '<Planta> <Servicio>' con servicio Energia o Gas "
    "(p. ej. 'Ensamble', 'RC Energia', 'RC GAS').",
    "Fila 1: notas libres (el pipeline la ignora). Fila 2: encabezados exactos de esta plantilla. "
    "Desde la fila 3: datos.",
    "El servicio se reconoce por los encabezados: columna 'KWH' = energía (kWh); columna '(m3)' = gas (m3).",
    "Un bloque de filas por mes, en orden cronológico. La última fila de cada mes es el total de la planta y "
    "su Proceso se llama igual que la planta (hoja 'RC GAS' -> Proceso 'RC').",
    "Energía: cada proceso lleva su Indice (participación); los índices de los procesos de un mes suman 1.",
    "Gas: una sola fila por mes (la fila total).",
    "Puede usar fórmulas, pero guarde el archivo desde Excel para que queden los valores calculados.",
    "Al cargar el archivo se REEMPLAZA el Excel completo: incluya todas las plantas y todos los meses.",
    "Para agregar una planta: agregue una hoja nueva con estas reglas. Se reentrenan los modelos de todas.",
    "Hojas que se ignoran: 'Instrucciones' y las que empiezan por '_' (como los ejemplos).",
    "Para usar un ejemplo: copie la hoja, quite el '_' del nombre y reemplace los datos.",
]

# Ejemplo: planta "Ejemplo", tres meses.
EJ_MESES = [(2026, "Ene", 100_000, (20_000, 15_000)), (2026, "Feb", 110_000, (22_000, 16_000)),
            (2026, "Mar", 95_000, (19_000, 14_500))]
EJ_INDICES = {"Linea 1": 0.6, "Linea 2": 0.4}
EJ_PPTO, EJ_COSTO = 2.9, 650


def _encabezado(ws, servicio):
    for j, h in enumerate(ENCABEZADOS[servicio], start=1):
        nota = ws.cell(row=1, column=j, value=DESCRIPCION[h][0] or None)
        nota.font = NOTA
        c = ws.cell(row=2, column=j, value=h)
        c.font, c.fill = BLANCO_NEGRITA, AZUL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = max(10, len(h) + 4)
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "A3"
    dv = DataValidation(type="list", formula1='"' + ",".join(MESES) + '"', errorStyle="warning",
                        error="Use Ene..Dic", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B3:B2000")


def _formato(encabezado):
    h = encabezado.lower()
    if "/equ" in h:
        return "0.000"
    if h in ("indice", "participacion"):
        return "0.00"
    if h in ("kwh", "kwh por proceso", "(m3)", "produccion (equivalente)") or "costo" in h:
        return "#,##0"
    return None


def _fila(ws, valores, total=False):
    ws.append(valores)
    for c in ws[ws.max_row]:
        fmt = _formato(ws.cell(row=2, column=c.column).value)
        if fmt:
            c.number_format = fmt
        if total:
            c.font, c.fill = NEGRITA, GRIS


def _ejemplo_energia(ws):
    _encabezado(ws, "energia")
    for anio, mes, kwh, prods in EJ_MESES:
        for (proc, idx), prod in zip(EJ_INDICES.items(), prods):
            consumo = kwh * idx
            _fila(ws, [anio, mes, "N/A", "N/A", proc, idx, kwh, consumo, prod, consumo / prod,
                       EJ_PPTO, EJ_COSTO, consumo * EJ_COSTO])
        prod = sum(prods)
        _fila(ws, [anio, mes, "N/A", "N/A", "Ejemplo", 1, kwh, kwh, prod, kwh / prod,
                   EJ_PPTO, EJ_COSTO, kwh * EJ_COSTO], total=True)


def _ejemplo_gas(ws):
    _encabezado(ws, "gas")
    for anio, mes, kwh, prods in EJ_MESES:
        m3, prod = kwh * 0.4, sum(prods)
        _fila(ws, [anio, mes, "N/A", "N/A", "Ejemplo", m3, prod, m3 / prod, 1.1, 2_100, m3 * 2_100],
              total=True)


def _instrucciones(ws):
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 90
    ws["A1"] = "Plantilla del Excel fuente — Dashboard energético Clarios"
    ws["A1"].font = Font(bold=True, size=14)
    fila = 3
    ws.cell(row=fila, column=1, value="Reglas").font = Font(bold=True, size=12)
    for i, r in enumerate(REGLAS, start=1):
        fila += 1
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=3)
        c = ws.cell(row=fila, column=1, value=f"{i}. {r}")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[fila].height = 30
    for servicio, titulo in [("energia", "Columnas — hoja de energía"), ("gas", "Columnas — hoja de gas")]:
        fila += 2
        ws.cell(row=fila, column=1, value=titulo).font = Font(bold=True, size=12)
        fila += 1
        for j, h in enumerate(["Encabezado (fila 2)", "Obligatoria", "Contenido"], start=1):
            c = ws.cell(row=fila, column=j, value=h)
            c.font, c.fill = BLANCO_NEGRITA, AZUL
        for h in ENCABEZADOS[servicio]:
            fila += 1
            _, obligatoria, desc = DESCRIPCION[h]
            ws.cell(row=fila, column=1, value=h)
            ws.cell(row=fila, column=2, value="Sí" if obligatoria else "No")
            ws.cell(row=fila, column=3, value=desc).alignment = Alignment(wrap_text=True, vertical="top")


def _impresion(ws):
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0


def crear(path=SALIDA):
    wb = Workbook()
    _instrucciones(wb.active)
    wb.active.title = "Instrucciones"
    _ejemplo_energia(wb.create_sheet("_Ejemplo Energia"))
    _ejemplo_gas(wb.create_sheet("_Ejemplo Gas"))
    for ws in wb.worksheets:
        _impresion(ws)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    from pathlib import Path
    print(crear(Path(sys.argv[1]) if len(sys.argv) > 1 else SALIDA))
